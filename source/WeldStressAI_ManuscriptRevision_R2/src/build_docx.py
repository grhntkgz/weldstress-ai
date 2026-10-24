from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=Path(r'E:\welding\outputs\WeldStressAI_ManuscriptRevision_R2')
source=(D/'MANUSCRIPT_REVISION_R2.md').read_text(encoding='utf-8');doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption','Header','Footer']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 for fonts in st.element.xpath('./w:rPr/w:rFonts'):
  for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+key),None)
 for borders in st.element.xpath('./w:pPr/w:pBdr'):borders.getparent().remove(borders)
 for color in st.element.xpath('./w:rPr/w:color'):
  for key in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+key),None)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].paragraph_format.space_after=Pt(14)
for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].paragraph_format.space_after=Pt(9);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False
header=sec.header.paragraphs[0];header.text='WeldStress AI   |   Scientific manuscript revision R2';header.style='Header';header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run();r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def keep(p,next=False):p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=next
def clean(s):return s.replace('**','').replace('`','')
def mathtext(p,s):
 om=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=s;mr.append(mt);om.append(mr);p._p.append(om)
def mr(s):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;r.append(t);return r
def sub(s,index='i'):
 n=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(s));b=OxmlElement('m:sub');b.append(mr(index));n.extend([e,b]);return n
def square(nodes):
 n=OxmlElement('m:sSup');e=OxmlElement('m:e');e.extend(nodes);s=OxmlElement('m:sup');s.append(mr('2'));n.extend([e,s]);return n
def summation(nodes,index='i'):
 n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');pr.append(ch);hide=OxmlElement('m:supHide');hide.set(qn('m:val'),'1');pr.append(hide);n.append(pr);b=OxmlElement('m:sub');b.append(mr(index));n.append(b);n.append(OxmlElement('m:sup'));e=OxmlElement('m:e');e.extend(nodes);n.append(e);return n
def fraction(top,bottom):
 n=OxmlElement('m:f');a=OxmlElement('m:num');a.extend(top);b=OxmlElement('m:den');b.extend(bottom);n.extend([a,b]);return n
def radical(node):
 n=OxmlElement('m:rad');pr=OxmlElement('m:radPr');h=OxmlElement('m:degHide');h.set(qn('m:val'),'1');pr.append(h);n.append(pr);n.append(OxmlElement('m:deg'));e=OxmlElement('m:e');e.append(node);n.append(e);return n
def equation(line):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;om=OxmlElement('m:oMath')
 if line.startswith('T2_i'):
  om.extend([sub('T2'),mr(' = '),fraction([summation([sub('a','j'),sub('s','j')],'j ∈ N(i)')],[summation([sub('a','j')],'j ∈ N(i)')]),mr('   (1)')])
 else:
  name=line.split(' =')[0];idx={'L2':2,'MAE':3,'RMSE':4}[name];lhs=sub('L','2') if name=='L2' else mr(name);error=[mr('('),sub('p'),mr(' − '),sub('y'),mr(')')]
  numerator=[summation([sub('w'),square(error)])] if name!='MAE' else [summation([sub('w'),mr('|'),sub('p'),mr(' − '),sub('y'),mr('|')])]
  denominator=[summation([sub('w'),square([sub('y')])])] if name=='L2' else [summation([sub('w')])]
  val=fraction(numerator,denominator);val=radical(val) if name!='MAE' else val;om.extend([lhs,mr(' = '),val,mr(f'   ({idx})')])
 p._p.append(om);keep(p)
lines=source.splitlines();i=0;paragraphs=0;figures=0;tables=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
 elif line.startswith('### '):doc.add_heading(line[4:],level=2)
 elif line.startswith('## '):doc.add_heading(line[3:],level=1)
 elif line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].lstrip().startswith('|'):
   vals=[clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[-: ]+',x or '-') for x in vals):rows.append(vals)
   i+=1
  table=doc.add_table(rows=0,cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False;cols=len(rows[0]);widths={5:[1.9,1.0,1.05,1.25,1.7],6:[.7,1.05,.8,1.0,1.5,1.85]}.get(cols,[6.9/cols]*cols)
  # Deliberately compact identifiers, generous metric headers.
  if rows[0][0]=='Model' and cols==6:widths=[.55,1.15,.95,.8,1.5,1.95]
  if rows[0][0]=='Held-out plate':widths=[1.0,1.475,1.475,1.475,1.475]
  for col,w in zip(table.columns,widths):col.width=Inches(w)
  for j,vals in enumerate(rows):
   cells=table.add_row().cells
   for k,(cell,value) in enumerate(zip(cells,vals)):
    cell.width=Inches(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders);margins=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEE' if j==0 else 'FFFFFF' if j%2 else 'F6F6F6');pr.append(shade)
    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if k==0 else WD_ALIGN_PARAGRAPH.CENTER;value='Context' if j==0 and value=='Observation' else value;r=p.add_run(value.replace('B2_GEOMETRY_XGBOOST','XGBoost').replace('B1_GEOMETRY_RIDGE','Ridge').replace('B0_CONSTANT','Constant'));r.font.size=Pt(9);r.bold=j==0
   trpr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
   if j==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
   group_next=(len(rows)<=7 and j<len(rows)-1) or (j<3) or (j>=len(rows)-3 and j<len(rows)-1)
   if group_next:
    for cell in cells:
     for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
  tables+=1;doc.add_paragraph().paragraph_format.space_after=Pt(2);continue
 elif line.startswith('!['):
  path=re.search(r'\]\((.+)\)',line).group(1);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(D/path),width=Inches(6.9));keep(p,True);figures+=1
 elif re.match(r'^(T2_i|L2 =|MAE =|RMSE =)',line):equation(line)
 elif re.match(r'^(Figure|Table) \d+\.',line):
  p=doc.add_paragraph(clean(line),style='Caption');keep(p,line.startswith('Table '))
 else:
  p=doc.add_paragraph(clean(line));paragraphs+=1
  if line.startswith('Keywords:'):p.runs[0].italic=True
  if line.endswith('The operator is') or line.endswith('Whole-field error is'):p.paragraph_format.keep_with_next=True
 i+=1
doc.core_properties.title='Geometry representation and observation context in learning stress fields from scanned welded joints';doc.core_properties.subject='Scientific manuscript revision from frozen numerical evidence';doc.core_properties.author='';doc.core_properties.keywords='weld geometry; numerical stress surrogate; exploratory plate-grouped study'
dest=D/'MANUSCRIPT_REVISION_R2.docx';doc.save(dest)
(D/'qa'/'DOCX_BUILD.json').write_text(json.dumps(dict(path=str(dest),figures=figures,tables=tables,body_paragraphs=paragraphs,page_size='Letter portrait',math='Editable native OMML linear equations',author_metadata='not invented'),indent=2))
print(json.dumps({'docx':str(dest),'figures':figures,'tables':tables}))
