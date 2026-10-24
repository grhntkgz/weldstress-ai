from common import *
from docx import Document
from pypdf import PdfReader
import zipfile
import xml.etree.ElementTree as ET

checks=[]
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math','a':'http://schemas.openxmlformats.org/drawingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
main=D/'05_Final_Manuscript/WELDSTRESS_AI_JCSR_VISUAL_FINAL.docx';si=D/'04_Supplementary_Update/SUPPLEMENTARY_INFORMATION.docx'
old=Document(J/'01_Main_Manuscript/WELDSTRESS_AI_JCSR_MAIN_MANUSCRIPT.docx');new=Document(main)
assert [[[c.text for c in r.cells] for r in t.rows] for t in old.tables]==[[[c.text for c in r.cells] for r in t.rows] for t in new.tables]
for p,count in [(main,9),(si,30)]:
 with zipfile.ZipFile(p) as z:
  xml=ET.fromstring(z.read('word/document.xml'));images=xml.findall('.//a:blip',ns);equations=xml.findall('.//m:oMath',ns);tables=xml.findall('.//w:tbl',ns)
  assert len(images)==count
  if p==main:assert len(equations)==4 and len(tables)==7
  hashes={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
  registry=rows(D/('05_Final_Manuscript/FINAL_FIGURE_REGISTRY.csv' if p==main else '04_Supplementary_Update/SUPPLEMENTARY_FIGURE_INDEX.csv'))
  assert all(r['SHA256'] in hashes for r in registry)
  assert not any(n.startswith('word/comments') for n in z.namelist())
  assert not xml.findall('.//w:ins',ns) and not xml.findall('.//w:del',ns)
  assert all(int(n.attrib['cx'])<=int(6.9*914400)+1 for n in xml.findall('.//wp:extent',ns))
  checks.append({'file':str(p),'images':len(images),'editable_tables':len(tables),'native_OMML_equations':len(equations),'all_embedded_figure_bytes_exact':True,'comments_tracked_changes':False,'widths_within_margins':True})
copy(D/'06_Scientific_QA/rendered/main/WELDSTRESS_AI_JCSR_VISUAL_FINAL.pdf','05_Final_Manuscript/WELDSTRESS_AI_JCSR_VISUAL_FINAL_REVIEW.pdf')
copy(D/'06_Scientific_QA/rendered/supplement_final/SUPPLEMENTARY_INFORMATION.pdf','04_Supplementary_Update/SUPPLEMENTARY_INFORMATION_REVIEW.pdf')
pages=[]
for kind,doc,pdf in [('main',main,D/'05_Final_Manuscript/WELDSTRESS_AI_JCSR_VISUAL_FINAL_REVIEW.pdf'),('supplement',si,D/'04_Supplementary_Update/SUPPLEMENTARY_INFORMATION_REVIEW.pdf')]:
 reader=PdfReader(pdf);alltext='\n'.join(p.extract_text() for p in reader.pages)
 captions=re.findall(r'Figure (S?\d+)\.',alltext)
 expected=[str(i) for i in range(1,10)] if kind=='main' else ['S'+str(i) for i in range(1,31)]
 assert set(expected).issubset(captions)
 for i,p in enumerate(reader.pages,1):
  text=p.extract_text();assert len(text.strip())>20
  pages.append({'document':kind,'page':i,'text_characters':len(text),'figure_captions':'; '.join(re.findall(r'Figure (S?\d+)\.',text)),'PDF_SHA256':sha(pdf),'DOCX_SHA256':sha(doc),'rendered_PNG':str(D/'06_Scientific_QA/rendered'/('main' if kind=='main' else 'supplement_final')/f'page-{i}.png')})
 assert '\ufffd' not in alltext
table('06_Scientific_QA/PAGE_RENDERING_INVENTORY.csv',pages)
js('06_Scientific_QA/DOCUMENT_STRUCTURE_AUDIT.json',{'status':'PASS','tables_identical_to_original_DOCX':True,'documents':checks,'main_pages':sum(r['document']=='main' for r in pages),'SI_pages':sum(r['document']=='supplement' for r in pages),'PDF_conversion':'Installed native Microsoft Word COM; packaged render_docx.py Poppler rasterization','author_placeholders_retained':all(x in '\n'.join(p.text for p in new.paragraphs) for x in ['AUTHOR CONFIRMATION REQUIRED','METHOD_IMPROVEMENT_MIXED'])})
print('DOCUMENT AUDIT PASS',len(pages),'pages')
